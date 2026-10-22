#!/usr/bin/env python3
"""Lit le tableur des armes et armures (Manuel des joueurs 2014, ch. 5) et produit data/equipment.json.

    python3 outils/equipement_xlsx.py                                   # outils/sources/Armes_et_ArmuresDnD.xlsx
    python3 outils/equipement_xlsx.py autre.xlsx --sortie /tmp/e.json

Feuille « Armes » : Arme, VO, Dégât, Poids, Prix, Propriétés ; les lignes de titre donnent la catégorie
(courante / de guerre, corps à corps / distance). Une ligne sans nom prolonge les propriétés de la précédente.
Feuille « Armures » : Armure, VO, CA, Force, Discrétion, Poids, Prix ; titres : légères, intermédiaires, lourdes, bouclier.
Ajouts hors tableur : armure métallique ou non (le druide refuse le métal), texte des propriétés « spécial ».
Prérequis : openpyxl.
"""
import argparse, json, os, re, unicodedata
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ap = argparse.ArgumentParser()
ap.add_argument("tableur", nargs="?", default=os.path.join(ROOT, "outils", "sources", "Armes_et_ArmuresDnD.xlsx"))
ap.add_argument("--sortie", default=os.path.join(ROOT, "data", "equipment.json"))
A = ap.parse_args()

def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")
def txt(v): return (str(v).replace("\xa0", " ").strip() if v is not None else "")

SPECIAL = {  # Manuel des joueurs, « Armes spéciales »
    "lance": "Désavantage contre une cible à 1,50 m ou moins. À deux mains sauf si tu es monté.",
    "net": "Une créature de taille Grande ou moins touchée est entravée jusqu'à sa libération (action : Force DD 10, ou 5 dégâts tranchants au filet, CA 10). Sans effet sur les créatures informes ou Très grandes. Une seule attaque par action, quel que soit le nombre d'attaques possibles.",
}
NOT_METAL = {"padded", "leather", "studded-leather", "hide"}

wb = openpyxl.load_workbook(A.tableur, data_only=True)
weapons, cat = [], None
for row in list(wb["Armes"].iter_rows(values_only=True))[1:]:
    n, vo, dmg, w, p, props = (txt(c) for c in row)
    if n and not vo:  # titre de section
        cat = ("simple" if "courantes" in n else "martial", "ranged" if "distance" in n else "melee"); continue
    if not n:  # suite des propriétés de la ligne précédente
        weapons[-1]["props"] = (weapons[-1]["props"].rstrip(", ") + ", " + props).strip(", "); continue
    m = re.match(r"(\d+(?:d\d+)?)\s+(\w+)", dmg)
    weapons.append({"id": slug(vo), "n": n, "vo": vo, "cat": cat[0], "rng": cat[1],
                    "dmg": m.group(1) if m else "", "dt": m.group(2) if m else "",
                    "props": "" if props == "-" else props, "w": "" if w == "-" else w, "p": p})
for x in weapons:  # propriétés lisibles par le code
    pr = x["props"].lower()
    f = {"finesse": "finesse" in pr, "light": "légère" in pr, "heavy": "lourde" in pr, "two": "deux mains" in pr,
         "reach": "allonge" in pr, "thrown": "lancer" in pr, "ammo": "munitions" in pr, "loading": "chargement" in pr}
    v = re.search(r"polyvalente \((\d+d\d+)\)", pr); r = re.search(r"portée ([\d,]+ ?m) ?/ ?([\d,]+ ?m)", x["props"])
    x.update({k: True for k, on in f.items() if on})
    if v: x["vers"] = v.group(1)
    if r: x["range"] = f"{r.group(1).replace(' ', '')}/{r.group(2).replace(' ', '')}".replace("m/", " m/").rstrip("m") + " m"
    if x["id"] in SPECIAL: x["special"] = SPECIAL[x["id"]]

armors, cat = [], None
CATS = {"légères": "light", "intermédiaires": "medium", "lourdes": "heavy", "Bouclier": "shield"}
for row in list(wb["Armures"].iter_rows(values_only=True))[1:]:
    n, vo, ac, st, stl, w, p = (txt(c) for c in row)
    if n and not vo: cat = next(v for k, v in CATS.items() if k in n); continue
    m = re.match(r"(\d+)", ac); dmax = re.search(r"max \+(\d)", ac); s = re.search(r"(\d+)", st)
    a = {"id": slug(vo), "n": n, "vo": vo, "cat": cat, "ac": int(m.group(1)), "w": w, "p": p}
    if "Dex" in ac: a["dex"] = True
    if dmax: a["dexMax"] = int(dmax.group(1))
    if s: a["str"] = int(s.group(1))
    if stl.startswith("Désavantage"): a["stealth"] = True
    if cat != "shield" and a["id"] not in NOT_METAL: a["metal"] = True
    armors.append(a)

assert len(weapons) == 37 and len(armors) == 13, (len(weapons), len(armors))
json.dump({"source": "Manuel des joueurs 2014, chapitre 5", "weapons": weapons, "armors": armors},
          open(A.sortie, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print(f"{len(weapons)} armes, {len(armors)} armures et bouclier → {A.sortie}")
